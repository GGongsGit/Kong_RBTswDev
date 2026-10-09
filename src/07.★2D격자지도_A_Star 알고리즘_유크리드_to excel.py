from heapq import heappush, heappop
from math import sqrt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def astar(grid, start, goal):
    grid = np.array(grid, dtype=int)
    R, C = grid.shape
    sr, sc = start
    gr, gc = goal
    
    NEI = [(1,0,1),(-1,0,1),(0,1,1),(0,-1,1),
           (1,1,sqrt(2)),(1,-1,sqrt(2)),(-1,1,sqrt(2)),(-1,-1,sqrt(2))]
    
    def h(r, c):
        return sqrt((r - gr)**2 + (c - gc)**2)
    
    open_heap = [(h(sr, sc), sr, sc)]
    g_score = {(sr, sc): 0}
    came_from = {}
    
    while open_heap:
        _, r, c = heappop(open_heap)
        
        if (r, c) == (gr, gc):
            path = []
            while (r, c) in came_from:
                path.append((r, c))
                r, c = came_from[(r, c)]
            path.append((sr, sc))
            return path[::-1], g_score[(gr, gc)]
        
        for dr, dc, cost in NEI:
            nr, nc = r + dr, c + dc
            if 0 <= nr < R and 0 <= nc < C and grid[nr, nc] == 0:
                new_g = g_score[(r, c)] + cost
                if (nr, nc) not in g_score or new_g < g_score[(nr, nc)]:
                    g_score[(nr, nc)] = new_g
                    came_from[(nr, nc)] = (r, c)
                    f = new_g + h(nr, nc)
                    heappush(open_heap, (f, nr, nc))
    
    return None, float('inf')

def save_to_excel(grid, start, goal, filename="astar_grid.xlsx"):
    path, cost = astar(grid, start, goal)
    R, C = len(grid), len(grid[0])
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Grid"
    
    thin_border = Border(
        left=Side(style='thin', color="000000"),
        right=Side(style='thin', color="000000"),
        top=Side(style='thin', color="000000"),
        bottom=Side(style='thin', color="000000")
    )
    
    # 그리드 그리기
    for r in range(R):
        ws.row_dimensions[r+1].height = 20
        for c in range(C):
            cell = ws.cell(row=r+1, column=c+1)
            ws.column_dimensions[get_column_letter(c+1)].width = 3
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')
            
            # 벽 = 검은색, 통로 = 흰색
            if grid[r][c] == 1:
                cell.fill = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
            else:
                cell.fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    # 경로 표시
    if path:
        for r, c in path:
            cell = ws.cell(row=r+1, column=c+1)
            
            if (r, c) == start:
                cell.fill = PatternFill(start_color="00B050", end_color="00B050", fill_type="solid")
                cell.value = "S"
                cell.font = Font(bold=True, color="FFFFFF")
            elif (r, c) == goal:
                cell.fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
                cell.value = "G"
                cell.font = Font(bold=True, color="FFFFFF")
            else:
                cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    
    wb.save(filename)
    print(f"✓ 저장: {filename}")
    print(f"✓ 비용: {cost:.3f}, 경로: {len(path) if path else 0} 노드")

if __name__ == "__main__":
    grid = [
        [0,0,0,1,0,1,0],
        [1,1,0,1,0,1,0],
        [0,0,0,0,0,1,0],
        [0,1,0,0,1,0,0],
        [0,1,1,0,1,1,0],
        [0,1,0,0,0,0,0],
        [0,0,0,1,0,1,0],
    ]
    
    save_to_excel(grid, (0, 0), (6, 6))
