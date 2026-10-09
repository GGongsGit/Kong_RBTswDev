# Astar_path.txt + Astar_block.txt → 경로 + Block 분산형 차트 (A_star 알고리즘 그래프 생성 방법_엑셀.xlsx 의 차트 1)
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties

PATH_FILE  = r"C:\KONG_2027Prg\RBTSW_Results\Astar_path.txt"
BLOCK_FILE = r"C:\KONG_2027Prg\RBTSW_Results\Astar_block.txt"
OUT_FILE   = r"C:\KONG_2027Prg\RBTSW_Results\astar_chart.xlsx"

def read_table(path):
    with open(path, "r", encoding="utf-8-sig") as f:
        head = f.readline().split()
        return head, [[float(v) for v in ln.split()] for ln in f if ln.strip()]

ph, path = read_table(PATH_FILE)       # Step 행 열 누적거리
bh, block = read_table(BLOCK_FILE)     # 열 행

wb = Workbook()
ws = wb.active
ws.title = "Astar"
ws.append(ph); [ws.append(r) for r in path]                  # A1:D(n+1)
ws["F1"], ws["G1"] = bh                                         # F1:G1 머리글
for i, r in enumerate(block, 2):
    ws.cell(i, 6, r[0]); ws.cell(i, 7, r[1])                   # F=열(X), G=행(Y)

n, m = len(path), len(block)
ch = ScatterChart()
ch.title = "A* 경로"
ch.style = 13

# 계열 1: Block  (선 없음, 검정 사각형 표식)
blk = Series(Reference(ws, min_col=7, min_row=2, max_row=m + 1),
             Reference(ws, min_col=6, min_row=2, max_row=m + 1), title="Block")
blk.marker = Marker(symbol="square", size=30)
blk.marker.graphicalProperties = GraphicalProperties(solidFill="000000")
blk.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
# 계열 2: Path   (X = 열 C열, Y = 행 B열)
pth = Series(Reference(ws, min_col=2, min_row=2, max_row=n + 1),
             Reference(ws, min_col=3, min_row=2, max_row=n + 1), title="Path")
pth.marker = Marker(symbol="circle", size=6)
pth.smooth = blk.smooth = False
ch.series += [blk, pth]                                         # Block 먼저 → Path가 위에 그려짐

for ax in (ch.x_axis, ch.y_axis):                               # 칸 경계(±0.5)와 중심(정수)에 눈금선
    ax.scaling.min, ax.scaling.max, ax.majorUnit = -0.5, 6.5, 0.5
    ax.delete = False
ch.x_axis.title, ch.y_axis.title = "열 (col)", "행 (row)"
ch.width = ch.height = 15
ws.add_chart(ch, "I2")
wb.save(OUT_FILE)
name = OUT_FILE.replace("\\", "/").split("/")[-1]       # 파일 이름만
print(f"저장: {name}  Path {n}점, Block {m}칸")
