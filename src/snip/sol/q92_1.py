# 10-2 예상문제 1 : 격자 TXT → 엑셀 수식(INDEX/MOD/INT)으로 Block 좌표 → 경로 + Block 차트
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties

GRID_FILE = r"C:\KONG_2027Prg\RBTSW_Results\q92_1_grid.txt"
PATH_FILE = r"C:\KONG_2027Prg\RBTSW_Results\q92_1_path.txt"
OUT_FILE  = r"C:\KONG_2027Prg\RBTSW_Results\q92_1_output.xlsx"

with open(GRID_FILE, "r", encoding="utf-8-sig") as f:
    grid = [[int(v) for v in ln.split()] for ln in f if ln.strip()]
with open(PATH_FILE, "r", encoding="utf-8-sig") as f:
    ph = f.readline().split()
    path = [[float(v) for v in ln.split()] for ln in f if ln.strip()]
R, C = len(grid), len(grid[0])

wb = Workbook(); ws = wb.active; ws.title = "Astar_grid"
for c in range(C):
    ws.cell(1, c + 1, f"C{c}")
for r in range(R):
    for c in range(C):
        ws.cell(r + 2, c + 1, grid[r][c])                      # A2:G8 = 격자 (2행 = 행 0)
last = chr(ord("A") + C - 1)
ws["J1"], ws["K1"] = "열(X)", "행(Y)"
for k in range(R * C):                                          # J2:K50 = 칸마다 1이면 좌표, 아니면 #N/A
    row = k + 2
    rr, cc = f"INT((ROW()-2)/{C})", f"MOD(ROW()-2,{C})"
    cell = f"INDEX($A$2:${last}${R+1},{rr}+1,{cc}+1)"
    ws.cell(row, 10, f"=IF({cell}=1,{cc},NA())")
    ws.cell(row, 11, f"=IF({cell}=1,{rr},NA())")
top = R + 4                                                     # 경로 표는 격자 아래
ws.cell(top, 1, ph[0]); ws.cell(top, 2, ph[1]); ws.cell(top, 3, ph[2]); ws.cell(top, 4, ph[3])
for i, p in enumerate(path, 1):
    for j, v in enumerate(p):
        ws.cell(top + i, j + 1, v)

ch = ScatterChart(); ch.title = "A* 경로"; ch.style = 13
blk = Series(Reference(ws, min_col=11, min_row=2, max_row=R*C + 1),
             Reference(ws, min_col=10, min_row=2, max_row=R*C + 1), title="Block")
blk.marker = Marker(symbol="square", size=30)
blk.marker.graphicalProperties = GraphicalProperties(solidFill="000000")
blk.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
pth = Series(Reference(ws, min_col=2, min_row=top + 1, max_row=top + len(path)),
             Reference(ws, min_col=3, min_row=top + 1, max_row=top + len(path)), title="Path")
pth.marker = Marker(symbol="circle", size=6)
blk.smooth = pth.smooth = False
ch.series += [blk, pth]
for ax, mx in ((ch.x_axis, C), (ch.y_axis, R)):
    ax.scaling.min, ax.scaling.max, ax.majorUnit = -0.5, mx - 0.5, 0.5
    ax.delete = False
ch.width = ch.height = 15
ws.add_chart(ch, "M2")
wb.save(OUT_FILE)

blocks = [(c, r) for r in range(R) for c in range(C) if grid[r][c] == 1]
print(f"격자 {R}x{C}, Block {len(blocks)}칸 (수식 J2:K{R*C+1}), 경로 {len(path)}점 (A{top+1}:D{top+len(path)})")
print("Block (열,행):", " ".join(f"({c},{r})" for c, r in blocks))
