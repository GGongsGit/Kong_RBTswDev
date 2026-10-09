# 10-1 예상문제 1 : 관절 각도 → 엑셀 수식으로 FK 계산 → 분산형 차트
import math
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker

IN_FILE  = r"C:\KONG_2027Prg\RBTSW_Results\q91_1_input.txt"
OUT_FILE = r"C:\KONG_2027Prg\RBTSW_Results\q91_1_output.xlsx"
L1, L2 = 0.85, 0.65

with open(IN_FILE, "r", encoding="utf-8-sig") as f:
    thetas = [[float(v) for v in ln.split()] for ln in f if ln.strip()]

wb = Workbook(); ws = wb.active; ws.title = "FK"
ws["A1"], ws["B1"], ws["A2"], ws["B2"] = "L1", L1, "L2", L2          # 링크 길이 입력 칸
ws.append([])
ws.append(["자세", "θ1(deg)", "θ2(deg)", "x0", "y0", "x1", "y1", "x2", "y2"])   # 4행 머리글
for i, (t1, t2) in enumerate(thetas):
    r = 5 + i
    ws.append([f"자세 {i+1}", t1, t2, 0, 0,
               f"=$B$1*COS(RADIANS(B{r}))",                         # x1 = L1·cos θ1
               f"=$B$1*SIN(RADIANS(B{r}))",                         # y1
               f"=F{r}+$B$2*COS(RADIANS(B{r}+C{r}))",               # x2 = x1 + L2·cos(θ1+θ2)
               f"=G{r}+$B$2*SIN(RADIANS(B{r}+C{r}))"])              # y2

# 차트용: 자세마다 3점(원점, 관절, 말단)을 세로로 늘어놓은 표 (K~L열 X/Y 쌍)
ch = ScatterChart(); ch.title = "2링크 자세"; ch.style = 13
ch.x_axis.title, ch.y_axis.title = "X (m)", "Y (m)"
for i in range(len(thetas)):
    r, top = 5 + i, 4 * i + 2
    ws.cell(top - 1, 11, f"자세 {i+1}")
    for k, (cx, cy) in enumerate([("D", "E"), ("F", "G"), ("H", "I")]):
        ws.cell(top + k, 11, f"={cx}{r}")
        ws.cell(top + k, 12, f"={cy}{r}")
    s = Series(Reference(ws, min_col=12, min_row=top, max_row=top + 2),
               Reference(ws, min_col=11, min_row=top, max_row=top + 2), title=f"자세 {i+1}")
    s.marker = Marker(symbol="circle", size=7); s.smooth = False
    ch.series.append(s)
for ax in (ch.x_axis, ch.y_axis):
    ax.scaling.min, ax.scaling.max, ax.majorUnit = -0.6, 1.6, 0.2
    ax.delete = False
ch.width = ch.height = 14
ws.add_chart(ch, "N2")
wb.save(OUT_FILE)

# 검산 (Python으로 같은 값 계산)
print("자세  x1      y1      x2      y2")
for i, (t1, t2) in enumerate(thetas):
    a, b = math.radians(t1), math.radians(t1 + t2)
    x1, y1 = L1*math.cos(a), L1*math.sin(a)
    print(f"{i+1}   {x1:.4f}  {y1:.4f}  {x1 + L2*math.cos(b):.4f}  {y1 + L2*math.sin(b):.4f}")
