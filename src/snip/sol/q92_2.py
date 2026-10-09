# 10-2 예상문제 2 : 격자 TXT → A*(8방향) 계산 → 경로 TXT + 엑셀 차트까지 한 번에
from heapq import heappush, heappop
from math import sqrt, hypot
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties

IN_FILE   = r"C:\KONG_2027Prg\RBTSW_Results\q92_2_input.txt"
PATH_TXT  = r"C:\KONG_2027Prg\RBTSW_Results\q92_2_output.txt"
OUT_XLSX  = r"C:\KONG_2027Prg\RBTSW_Results\q92_2_output.xlsx"
START, GOAL = (0, 0), (6, 4)
NEI = [(1,0,1), (-1,0,1), (0,1,1), (0,-1,1),
       (1,1,sqrt(2)), (1,-1,sqrt(2)), (-1,1,sqrt(2)), (-1,-1,sqrt(2))]

with open(IN_FILE, "r", encoding="utf-8-sig") as f:
    grid = [[int(v) for v in ln.split()] for ln in f if ln.strip()]
R, C = len(grid), len(grid[0])

# ---- A* ----
h = lambda p: hypot(p[0] - GOAL[0], p[1] - GOAL[1])
g, came, pq = {START: 0.0}, {}, [(h(START), START)]
while pq:
    _, cur = heappop(pq)
    if cur == GOAL:
        break
    for dr, dc, w in NEI:
        n = (cur[0] + dr, cur[1] + dc)
        if 0 <= n[0] < R and 0 <= n[1] < C and grid[n[0]][n[1]] == 0 and g[cur] + w < g.get(n, 1e9):
            g[n], came[n] = g[cur] + w, cur
            heappush(pq, (g[n] + h(n), n))
path = [GOAL]
while path[-1] in came:
    path.append(came[path[-1]])
path.reverse()

# ---- TXT ----
with open(PATH_TXT, "w", encoding="utf-8") as f:
    f.write("Step 행 열 누적거리\n")
    acc = 0.0
    for i, (r, c) in enumerate(path):
        acc += hypot(r - path[i-1][0], c - path[i-1][1]) if i else 0.0
        f.write(f"{i+1} {r} {c} {acc:.3f}\n")

# ---- 엑셀 : 경로 표(A:D) + Block 목록(F:G) + 차트 ----
wb = Workbook(); ws = wb.active; ws.title = "Astar"
ws.append(["Step", "행", "열", "누적거리"])
acc = 0.0
for i, (r, c) in enumerate(path):
    acc += hypot(r - path[i-1][0], c - path[i-1][1]) if i else 0.0
    ws.append([i + 1, r, c, round(acc, 3)])
blocks = [(c, r) for c in range(C) for r in range(R) if grid[r][c] == 1]
ws["F1"], ws["G1"] = "열", "행"
for i, (c, r) in enumerate(blocks, 2):
    ws.cell(i, 6, c); ws.cell(i, 7, r)

ch = ScatterChart(); ch.title = f"A* 경로 (cost {g[GOAL]:.3f})"; ch.style = 13
blk = Series(Reference(ws, min_col=7, min_row=2, max_row=len(blocks) + 1),
             Reference(ws, min_col=6, min_row=2, max_row=len(blocks) + 1), title="Block")
blk.marker = Marker(symbol="square", size=30)
blk.marker.graphicalProperties = GraphicalProperties(solidFill="000000")
blk.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
pth = Series(Reference(ws, min_col=2, min_row=2, max_row=len(path) + 1),
             Reference(ws, min_col=3, min_row=2, max_row=len(path) + 1), title="Path")
pth.marker = Marker(symbol="circle", size=6)
blk.smooth = pth.smooth = False
ch.series += [blk, pth]
for ax, mx in ((ch.x_axis, C), (ch.y_axis, R)):
    ax.scaling.min, ax.scaling.max, ax.majorUnit = -0.5, mx - 0.5, 0.5
    ax.delete = False
ch.width = ch.height = 15
ws.add_chart(ch, "I2")
wb.save(OUT_XLSX)
print(f"경로 {len(path)}점, cost {g[GOAL]:.3f}, Block {len(blocks)}칸 → TXT, XLSX 저장")
