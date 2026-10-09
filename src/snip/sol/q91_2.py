# 10-1 예상문제 2 : 궤적 CSV → 시간-위치/속도/가속도 차트 (계열 3개)
import csv
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series

IN_FILE  = r"C:\KONG_2027Prg\RBTSW_Results\q91_2_input.csv"
OUT_FILE = r"C:\KONG_2027Prg\RBTSW_Results\q91_2_output.xlsx"

wb = Workbook(); ws = wb.active; ws.title = "BBPB"
with open(IN_FILE, "r", newline="", encoding="utf-8-sig") as f:
    r = csv.reader(f)
    ws.append(next(r))                                    # 머리글 그대로
    for row in r:
        ws.append([float(v) for v in row])                # 숫자로 넣어야 차트가 그려짐
n = ws.max_row

ch = ScatterChart(); ch.title = "BBPB 궤적"; ch.style = 13
ch.x_axis.title, ch.y_axis.title = "Time (s)", "값"
x = Reference(ws, min_col=1, min_row=2, max_row=n)
for col in (2, 3, 4):                                    # Pos, Vel, Acc
    s = Series(Reference(ws, min_col=col, min_row=1, max_row=n), x, title_from_data=True)
    s.marker.symbol = "none"                             # 점 1000개 → 표식 없이 선만
    s.smooth = False
    ch.series.append(s)
ch.x_axis.delete = ch.y_axis.delete = False
ch.width, ch.height = 20, 11
ws.add_chart(ch, "F2")
wb.save(OUT_FILE)

vals = list(ws.iter_rows(min_row=2, values_only=True))
print(f"데이터 {n-1}행, 시간 {vals[0][0]} ~ {vals[-1][0]} s")
for i, name in enumerate(["Pos", "Vel", "Acc"], 1):
    col = [v[i] for v in vals]
    print(f"{name}: min {min(col):.4f}, max {max(col):.4f}")
