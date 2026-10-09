# PxPy_data.txt → 엑셀 표 + 분산형 차트 (sample_data/robot_joint_positions.xlsx 와 같은 모양)
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker

IN_FILE  = r"C:\KONG_2027Prg\RBTSW_Results\PxPy_data.txt"
OUT_FILE = r"C:\KONG_2027Prg\RBTSW_Results\robot_joint_chart.xlsx"

# 1) TXT 읽기 : 머리글 1줄 + 자세별 6개 값
with open(IN_FILE, "r", encoding="utf-8-sig") as f:
    header = f.readline().split()
    rows = [[float(v) for v in line.split()] for line in f if line.strip()]

wb = Workbook()
ws = wb.active
ws.title = "관절위치"

# 2) 원본 표 (A1:G4)
ws.append(["자세"] + header)
for i, r in enumerate(rows, 1):
    ws.append([f"자세 {i}"] + r)

# 3) 차트용 표 (A6:G9) : 자세마다 X열, Y열  ← 분산형 차트는 '한 계열 = X 범위 + Y 범위'
ws["A6"] = "점"
for i in range(len(rows)):
    ws.cell(6, 2 + 2*i, f"자세{i+1} X")
    ws.cell(6, 3 + 2*i, f"자세{i+1} Y")
for k, name in enumerate(["P0 베이스", "P1 관절1", "P2 끝단"]):
    ws.cell(7 + k, 1, name)
    for i, r in enumerate(rows):
        ws.cell(7 + k, 2 + 2*i, r[2*k])        # Px_k
        ws.cell(7 + k, 3 + 2*i, r[2*k + 1])    # Py_k

# 4) 분산형(직선 및 표식) 차트
ch = ScatterChart()
ch.title = "로봇 관절 위치 (2링크)"
ch.style = 13
ch.x_axis.title = "X (m)"
ch.y_axis.title = "Y (m)"
for i, sym in enumerate(["circle", "square", "triangle"][:len(rows)]):
    xs = Reference(ws, min_col=2 + 2*i, min_row=7, max_row=9)
    ys = Reference(ws, min_col=3 + 2*i, min_row=7, max_row=9)
    s = Series(ys, xs, title=f"자세 {i+1}")
    s.marker = Marker(symbol=sym, size=8)
    s.smooth = False
    ch.series.append(s)
ch.x_axis.scaling.min, ch.x_axis.scaling.max, ch.x_axis.majorUnit = -0.4, 1.2, 0.2
ch.y_axis.scaling.min, ch.y_axis.scaling.max, ch.y_axis.majorUnit = -0.2, 1.4, 0.2
ch.x_axis.delete = ch.y_axis.delete = False          # 축 숫자 표시 (openpyxl 3.1 기본값은 숨김)
ch.width, ch.height = 16, 15                          # cm, 정사각형에 가깝게
ws.add_chart(ch, "I1")

wb.save(OUT_FILE)
name = OUT_FILE.replace("\\", "/").split("/")[-1]       # 파일 이름만
print(f"저장: {name}  계열 {len(ch.series)}개, X축 -0.4~1.2, Y축 -0.2~1.4")
